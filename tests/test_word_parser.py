from __future__ import annotations

import subprocess
import sys
from typing import TYPE_CHECKING, Any, cast

import pytest
from docx import Document

from venture_agents.tools.word_parser import NumberingDefinitions, parse_word_to_markdown


if TYPE_CHECKING:
    from pathlib import Path


def test_parse_word_to_markdown_preserves_headings_and_lists(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Investment Memo", level=1)
    document.add_paragraph("Company overview.")
    document.add_heading("Highlights", level=2)
    document.add_paragraph("Large market", style="List Bullet")
    document.add_paragraph("Strong team", style="List Bullet")
    document.add_heading("Process", level=2)
    document.add_paragraph("Review materials", style="List Number")
    document.add_paragraph("Interview customers", style="List Number")

    docx_path = tmp_path / "memo.docx"
    document.save(str(docx_path))

    markdown = parse_word_to_markdown(docx_path)

    assert markdown == (
        "# Investment Memo\n\n"
        "Company overview.\n\n"
        "## Highlights\n\n"
        "- Large market\n"
        "- Strong team\n\n"
        "## Process\n\n"
        "1. Review materials\n"
        "2. Interview customers"
    )


def test_parse_word_to_markdown_preserves_nested_list_indentation(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Parent bullet", style="List Bullet")
    document.add_paragraph("Child bullet", style="List Bullet 2")
    document.add_paragraph("Second parent", style="List Bullet")

    docx_path = tmp_path / "nested.docx"
    document.save(str(docx_path))

    markdown = parse_word_to_markdown(docx_path)

    assert markdown == "- Parent bullet\n  - Child bullet\n- Second parent"


def test_parse_word_to_markdown_renders_tables_in_body_order(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Metrics", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Note"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "A|B"
    document.add_paragraph("After table.")

    docx_path = tmp_path / "table.docx"
    document.save(str(docx_path))

    markdown = parse_word_to_markdown(docx_path)

    assert markdown == ("# Metrics\n\n" "| Name | Note |\n" "| --- | --- |\n" "| Revenue | A\\|B |\n\n" "After table.")


def test_parse_word_to_markdown_rejects_non_docx_files(tmp_path: Path) -> None:
    file_path = tmp_path / "legacy.doc"
    file_path.write_text("not a docx file", encoding="utf-8")

    with pytest.raises(ValueError, match=r"Only \.docx"):
        parse_word_to_markdown(file_path)


def test_parse_word_to_markdown_rejects_missing_files(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="Word document not found"):
        parse_word_to_markdown(tmp_path / "missing.docx")


def test_numbering_definitions_allows_documents_without_numbering_part() -> None:
    class Part:
        @property
        def numbering_part(self):
            raise NotImplementedError

    class DocumentWithoutNumberingPart:
        part = Part()

    NumberingDefinitions(cast("Any", DocumentWithoutNumberingPart()))


def test_word_parser_cli_writes_markdown_to_output_file(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("CLI Word", level=1)
    document.add_paragraph("Rendered from the command line.")

    docx_path = tmp_path / "cli.docx"
    output_path = tmp_path / "output.md"
    document.save(str(docx_path))

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "venture_agents.tools.word_parser",
            str(docx_path),
            "--output",
            str(output_path),
        ],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0
    assert result.stderr == ""
    assert result.stdout == ""
    assert output_path.read_text(encoding="utf-8") == "# CLI Word\n\nRendered from the command line.\n"
