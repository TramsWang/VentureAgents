from __future__ import annotations

import subprocess
import sys
from datetime import date
from typing import TYPE_CHECKING

from openpyxl import Workbook

from venture_agents.tools.excel_parser import SheetMd, parse_excel_to_markdown


if TYPE_CHECKING:
    from pathlib import Path


def test_parse_excel_to_markdown_returns_non_empty_sheets(tmp_path: Path) -> None:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"
    data_sheet.append(["Name", "Amount", "Note"])
    data_sheet.append(["Alpha", 1, "A|B"])
    data_sheet.append(["Beta", None, "line1\nline2"])

    empty_sheet = workbook.create_sheet("Empty")
    empty_sheet.append([None, None])

    whitespace_sheet = workbook.create_sheet("Whitespace")
    whitespace_sheet.append(["  ", None])

    workbook_path = tmp_path / "sample.xlsx"
    workbook.save(workbook_path)

    sheets = parse_excel_to_markdown(workbook_path)

    assert sheets == [
        SheetMd(
            sheet_name="Data",
            markdown=(
                "| Name | Amount | Note |\n| --- | --- | --- |\n| Alpha | 1 | A\\|B |\n| Beta |  | line1<br>line2 |"
            ),
        )
    ]


def test_parse_excel_to_markdown_trims_empty_edges_and_formats_dates(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet["B2"] = "Date"
    sheet["C2"] = "Value"
    sheet["B3"] = date(2026, 5, 11)
    sheet["C3"] = 0

    workbook_path = tmp_path / "sparse.xlsx"
    workbook.save(workbook_path)

    sheets = parse_excel_to_markdown(workbook_path)

    assert sheets == [
        SheetMd(
            sheet_name="Sparse",
            markdown="| Date | Value |\n| --- | --- |\n| 2026-05-11 | 0 |",
        )
    ]


def test_excel_parser_cli_outputs_markdown(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CLI"
    sheet.append(["Name", "Value"])
    sheet.append(["Gamma", 3])

    workbook_path = tmp_path / "cli.xlsx"
    workbook.save(workbook_path)

    result = subprocess.run(
        [sys.executable, "-m", "venture_agents.tools.excel_parser", str(workbook_path)],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0
    assert result.stderr == ""
    assert result.stdout == "## CLI\n\n| Name | Value |\n| --- | --- |\n| Gamma | 3 |\n"


def test_excel_parser_cli_writes_markdown_to_output_file(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output"
    sheet.append(["Name", "Value"])
    sheet.append(["Delta", 4])

    workbook_path = tmp_path / "output.xlsx"
    output_path = tmp_path / "output.md"
    workbook.save(workbook_path)

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "venture_agents.tools.excel_parser",
            str(workbook_path),
            "--output",
            str(output_path),
        ],
        capture_output=True,
        check=False,
        text=True,
    )

    assert result.returncode == 0
    assert result.stderr == ""
    assert result.stdout == ""
    assert output_path.read_text(encoding="utf-8") == "## Output\n\n| Name | Value |\n| --- | --- |\n| Delta | 4 |\n"
